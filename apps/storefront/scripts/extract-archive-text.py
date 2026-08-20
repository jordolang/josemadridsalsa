#!/usr/bin/env python3
"""Full-text extraction for the document archive's searchable index.

Walks every file under the local (gitignored) `Documents/` archive and pulls the
text layer out of the text-bearing formats (PDF, Word, spreadsheets, CSV/TXT,
HTML). Image-only scans have no text layer and are flagged `needsOcr` rather than
guessed at. Emits one JSONL record per file, keyed by its archive-relative path,
for `scripts/import-archive-documents.ts` to merge with the metadata index.

Usage:
    python3 scripts/extract-archive-text.py --archive ../../Documents --out <dir>

Output: <out>/archive-text.jsonl
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

from _spreadsheet import read_sheets

try:
    import docx  # python-docx
except ImportError:
    docx = None

MAX_CHARS = 300_000  # generous cap; almost every file is far smaller

TEXT_EXTS = {".txt", ".csv", ".md", ".log", ".xml", ".json"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".tif", ".tiff", ".heic", ".psd"}
SKIP_EXTS = {".zip", ".mov", ".mp4", ".lnk", ".url", ".aae", ".ps1", ".psm1", ".psd1", ".css", ".ds_store"}


def clip(text: str) -> str:
    text = text.replace("\x00", " ").strip()
    return text[:MAX_CHARS]


def strip_html(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    raw = re.sub(r"&nbsp;", " ", raw)
    return re.sub(r"\s+", " ", raw)


def extract_pdf(path: Path) -> str:
    try:
        out = subprocess.run(
            ["pdftotext", "-q", "-enc", "UTF-8", str(path), "-"],
            capture_output=True,
            timeout=60,
        )
        return out.stdout.decode("utf-8", "replace")
    except Exception:
        return ""


def extract_docx(path: Path) -> str:
    if docx is None:
        return extract_via_textutil(path)
    try:
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for tr in table.rows:
                parts.append("\t".join(c.text for c in tr.cells))
        return "\n".join(parts)
    except Exception:
        return extract_via_textutil(path)


def extract_via_textutil(path: Path) -> str:
    """macOS textutil handles .doc/.rtf and is a fallback for .docx."""
    try:
        out = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(path)],
            capture_output=True,
            timeout=60,
        )
        return out.stdout.decode("utf-8", "replace")
    except Exception:
        return ""


def extract_spreadsheet(path: Path) -> str:
    try:
        chunks: list[str] = []
        for name, rows in read_sheets(path):
            chunks.append(f"# {name}")
            for row in rows:
                cells = [str(c) for c in row if c not in (None, "")]
                if cells:
                    chunks.append("\t".join(cells))
                if sum(len(c) for c in chunks) > MAX_CHARS:
                    break
        return "\n".join(chunks)
    except Exception:
        return ""


def extract_text_file(path: Path) -> str:
    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return ""
    if path.suffix.lower() in (".html", ".htm", ".xml"):
        return strip_html(raw)
    return raw


def extract(path: Path) -> tuple[str, str, bool]:
    """Return (text, status, needs_ocr)."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = extract_pdf(path)
        if text.strip():
            return clip(text), "EXTRACTED", False
        return "", "SCANNED_NO_TEXT", True
    if ext == ".docx":
        return clip(extract_docx(path)), "EXTRACTED", False
    if ext in (".doc", ".rtf"):
        text = extract_via_textutil(path)
        return (clip(text), "EXTRACTED", False) if text.strip() else ("", "EMPTY", False)
    if ext in (".xlsx", ".xls"):
        text = extract_spreadsheet(path)
        return (clip(text), "EXTRACTED", False) if text.strip() else ("", "UNSUPPORTED", False)
    if ext in (".html", ".htm"):
        return clip(extract_text_file(path)), "EXTRACTED", False
    if ext in TEXT_EXTS:
        text = extract_text_file(path)
        return (clip(text), "EXTRACTED", False) if text.strip() else ("", "EMPTY", False)
    if ext in IMAGE_EXTS:
        return "", "SCANNED_NO_TEXT", True
    return "", "UNSUPPORTED", False


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--archive", default="../../Documents")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    archive_root = Path(args.archive).resolve()
    if not archive_root.is_dir():
        sys.exit(f"archive not found: {archive_root}")

    out_dir = Path(args.out).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / "archive-text.jsonl"

    files = sorted(
        p for p in archive_root.rglob("*")
        if p.is_file() and not p.name.startswith(".")
    )
    counts: dict[str, int] = {}
    done = 0
    with out_file.open("w", encoding="utf-8") as fh:
        for path in files:
            rel = str(path.relative_to(archive_root))
            # Skip the archive's own bookkeeping files.
            if rel in ("search_index.csv", "_move_manifest.csv", "README.md", "SEARCH.html"):
                continue
            if path.suffix.lower() in SKIP_EXTS:
                text, status, needs_ocr = "", "UNSUPPORTED", False
            else:
                text, status, needs_ocr = extract(path)
            counts[status] = counts.get(status, 0) + 1
            fh.write(
                json.dumps(
                    {
                        "path": rel,
                        "ext": path.suffix.lower().lstrip("."),
                        "textChars": len(text),
                        "extraction": status,
                        "needsOcr": needs_ocr,
                        "extractedText": text or None,
                    },
                    ensure_ascii=False,
                )
                + "\n"
            )
            done += 1
            if done % 250 == 0:
                print(f"  {done} files... {counts}", file=sys.stderr)

    print(f"\nextracted {done} files -> {out_file}")
    print(f"status: {counts}")


if __name__ == "__main__":
    main()
