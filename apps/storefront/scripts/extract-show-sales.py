#!/usr/bin/env python3
"""Stage 1 of the show-sales import.

Auto-detects the two sales-sheet shapes under `04 Shows & Events` and emits one
JSONL record per event day:

  SHOW   — header "Show Name" + "Sales": one row per show, with a "Show Person".
  MARKET — header "Farmers Market" + "Sales": rows grouped under a market name in
           the first column (carried down), with per-day date/sales/paid/expenses.

Totals rows and blank rows are dropped. Date parsing of ranges like "1/3-5/2025"
is left to the tested TS normalizer.

Usage: python3 scripts/extract-show-sales.py --archive ../../Documents --out <dir>
Output: <out>/show-sales.raw.jsonl
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl required")

SUBPATH = "04 Shows & Events"


def md5(path: Path) -> str:
    h = hashlib.md5()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def iso(v) -> str | None:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    return None


def year_from_path(rel: str) -> int | None:
    for seg in rel.split(os.sep):
        if re.fullmatch(r"(19|20)\d{2}", seg.strip()):
            return int(seg.strip())
    return None


def is_total(v) -> bool:
    return isinstance(v, str) and "total" in v.strip().lower()


def find_header(rows, *needles):
    for i, r in enumerate(rows[:6]):
        cells = [str(c).strip().lower() for c in r if c is not None]
        if all(any(n in c for c in cells) for n in needles):
            return i, [str(c).strip() if c is not None else "" for c in r]
    return None, None


def col_index(header, *names):
    low = [h.lower() for h in header]
    for n in names:
        for i, h in enumerate(low):
            if n in h:
                return i
    return None


def cell(r, i):
    return r[i] if i is not None and i < len(r) else None


def extract_show(rows, rel, digest, year) -> list[dict]:
    hi, header = find_header(rows, "show name", "sales")
    if hi is None:
        return []
    c_date = col_index(header, "show date", "date")
    c_name = col_index(header, "show name")
    c_sales = col_index(header, "sales")
    c_person = col_index(header, "show person", "person")
    out = []
    for offset, r in enumerate(rows[hi + 1 :], start=hi + 2):
        name = cell(r, c_name)
        if not name or is_total(name) or is_total(cell(r, 0)):
            continue
        dv = cell(r, c_date)
        out.append({
            "sourceFile": rel, "sourceMd5": digest, "sourceRow": offset,
            "year": year, "eventType": "SHOW",
            "showName": str(name).strip(),
            "dateText": None if iso(dv) else (str(dv).strip() if dv else None),
            "dateIso": iso(dv),
            "sales": cell(r, c_sales),
            "person": str(cell(r, c_person)).strip() if cell(r, c_person) else None,
        })
    return out


def extract_market(rows, rel, digest, year) -> list[dict]:
    hi, header = find_header(rows, "farmers mark", "sales")
    if hi is None:
        return []
    c_date = col_index(header, "date")
    c_sales = col_index(header, "sales")
    c_person = col_index(header, "sales person", "person")
    c_paid = col_index(header, "amount paid", "paid")
    c_exp = col_index(header, "exspense", "expense")
    current = None
    out = []
    for offset, r in enumerate(rows[hi + 1 :], start=hi + 2):
        first = cell(r, 0)
        if isinstance(first, str) and first.strip() and not is_total(first):
            current = first.strip()
        if is_total(first):
            continue
        dv = cell(r, c_date)
        sales = cell(r, c_sales)
        if dv is None and sales is None:
            continue
        if not current:
            continue
        out.append({
            "sourceFile": rel, "sourceMd5": digest, "sourceRow": offset,
            "year": year, "eventType": "FARMERS_MARKET",
            "showName": current,
            "dateText": None if iso(dv) else (str(dv).strip() if dv else None),
            "dateIso": iso(dv),
            "sales": sales,
            "amountPaid": cell(r, c_paid),
            "expenses": cell(r, c_exp),
            "person": str(cell(r, c_person)).strip() if cell(r, c_person) else None,
        })
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--archive", default="../../Documents")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    base = Path(args.archive).resolve() / SUBPATH
    if not base.is_dir():
        sys.exit(f"not found: {base}")
    archive_root = Path(args.archive).resolve()

    files = sorted(
        p for p in base.rglob("*")
        if p.is_file() and p.suffix.lower() in (".xlsx", ".xls") and not p.name.startswith("~")
    )
    out_dir = Path(args.out).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / "show-sales.raw.jsonl"

    total = 0
    with out_file.open("w", encoding="utf-8") as fh:
        for path in files:
            rel = str(path.relative_to(archive_root))
            try:
                wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
                rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
            except Exception:
                continue
            digest = md5(path)
            year = year_from_path(rel)
            recs = extract_show(rows, rel, digest, year) or extract_market(rows, rel, digest, year)
            for rec in recs:
                fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
            if recs:
                total += len(recs)
                print(f"  {len(recs):4d} rows  {rel}")

    print(f"\nwrote {total} rows -> {out_file}")


if __name__ == "__main__":
    main()
